"""The result as a spreadsheet: one sheet per record kind, every number with its interval."""
import json

from openpyxl import load_workbook
from synth import box_room, merge

from roomscan.export.build import build_output
from roomscan.export.sheet import write_sheet
from roomscan.geometry.layout import extract_layout


def test_workbook_has_every_room_wall_and_interval(tmp_path):
    a = box_room(4.0, 3.0, 2.7, origin=(0, 0), door=(1, 1.0, 0.9, 2.05))
    b = box_room(1.6, 2.0, 2.7, origin=(4.12, 0), door=(3, 1.1, 0.9, 2.05), seed=1)
    lay = extract_layout(merge(a, b))
    info = {"id": "synthetic", "tier": "lidar", "source": "test", "n_frames_used": 0, "meta": {}}
    res = json.loads(build_output(lay, [], "lidar", info, {"enabled": False, "method": "none"}).model_dump_json())
    wb = load_workbook(write_sheet(res, tmp_path / "result.xlsx"))
    assert wb.sheetnames == ["Summary", "Rooms", "Walls", "Openings", "Damage", "Flags", "Scope", "Warnings"]
    rooms, walls = wb["Rooms"], wb["Walls"]
    assert rooms.max_row - 1 == len(res["rooms"]) == 2
    assert walls.max_row - 1 == sum(len(r["walls"]) for r in res["rooms"])
    head = [c.value for c in walls[1]]
    i = head.index("length (m)")
    for row in walls.iter_rows(min_row=2, values_only=True):
        value, lo, hi = row[i:i + 3]
        assert lo <= value <= hi  # value with its 90 % interval
    areas = sorted(row[2] for row in rooms.iter_rows(min_row=2, values_only=True))
    assert abs(areas[1] - 12.0) < 0.1 and abs(areas[0] - 3.2) < 0.1
