"""The output contract as a spreadsheet (.xlsx): one sheet per kind of record.

Sheets: Summary, Rooms, Walls, Openings, Damage, Flags, Scope, Warnings. Every measured
number is written as three columns: value, 90 % interval low, 90 % interval high (a missing
measurement is left blank, with its lower bound where one exists). Units are metres and
square metres. Built from the published result.json, so it can be made for any run.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DDE6F0")


def _m(m: dict | None) -> list:
    """value, 90 % low, 90 % high of a measurement (blanks when not measured)."""
    if not m or m.get("value") is None:
        return [None, (m or {}).get("lower_bound"), None]
    lo, hi = m["ci90"] if m.get("ci90") else (None, None)
    r = lambda v: round(float(v), 4) if v is not None else None  # noqa: E731
    return [r(m["value"]), r(lo), r(hi)]


def _mcols(name: str, unit: str) -> list[str]:
    return [f"{name} ({unit})", f"{name} low", f"{name} high"]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), HEAD_FILL, Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)


def build_workbook(res: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    cap, prop = res.get("capture", {}), res.get("property", {})
    room_name = {r["id"]: r["id"] for r in res.get("rooms", [])}
    summary = [["capture", cap.get("id")], ["tier", cap.get("tier")], ["source", cap.get("source")],
               ["rooms", len(res.get("rooms", []))],
               ["footprint (m2)", *_m(prop.get("footprint_area"))],
               ["damage regions", len(res.get("damage", []))],
               ["concealed-damage flags", len(res.get("concealed_damage_flags", []))],
               ["scope items", len(res.get("scope", []))],
               ["interval", res.get("interval", "90 % interval: low - high")],
               ["schema version", res.get("schema_version")]]
    _sheet(wb, "Summary", ["item", "value", "90 % low", "90 % high"], summary)

    rows = []
    for r in res.get("rooms", []):
        rows.append([r["id"], r.get("label"), *_m(r.get("floor_area")), *_m(r.get("ceiling_height")),
                     *_m(r.get("perimeter")), len(r.get("walls", [])), len(r.get("openings", [])),
                     r.get("ceiling_source")])
    _sheet(wb, "Rooms", ["room", "label", *_mcols("floor area", "m2"), *_mcols("ceiling height", "m"),
                         *_mcols("perimeter", "m"), "walls", "openings", "ceiling source"], rows)

    rows = []
    for r in res.get("rooms", []):
        for w in r.get("walls", []):
            rows.append([r["id"], w["id"], *_m(w.get("length")), *_m(w.get("height")), *_m(w.get("area")),
                         round(float(w.get("evidence_coverage") or 0), 3), ", ".join(w.get("opening_ids") or [])])
    _sheet(wb, "Walls", ["room", "wall", *_mcols("length", "m"), *_mcols("height", "m"), *_mcols("area", "m2"),
                         "evidence coverage (0-1)", "openings"], rows)

    rows = []
    for r in res.get("rooms", []):
        for o in r.get("openings", []):
            rows.append([r["id"], o["id"], o.get("type"), o.get("wall_id"), o.get("connects_to"),
                         *_m(o.get("width")), *_m(o.get("height")), *_m(o.get("sill_height"))])
    _sheet(wb, "Openings", ["room", "opening", "type", "wall", "connects to", *_mcols("width", "m"),
                            *_mcols("height", "m"), *_mcols("sill height", "m")], rows)

    rows = [[d["id"], room_name.get(d.get("room_id"), d.get("room_id")), d.get("surface_id"), d.get("damage_class"),
             d.get("score"), *_m(d.get("area")), *_m(d.get("extent_u")), *_m(d.get("extent_v")),
             (d.get("center") or [None, None, None])[2], ", ".join(str(f) for f in d.get("evidence_frames", []))]
            for d in res.get("damage", [])]
    _sheet(wb, "Damage", ["damage", "room", "surface", "class", "score", *_mcols("area", "m2"),
                          *_mcols("width", "m"), *_mcols("height", "m"), "centre above floor (m)",
                          "evidence frames"], rows)

    rows = [[f["id"], f.get("room_id"), f.get("surface_id"), f.get("rule_id"), f.get("risk"), f.get("rule"),
             ", ".join(f.get("triggered_by", [])), f.get("recommendation")]
            for f in res.get("concealed_damage_flags", [])]
    _sheet(wb, "Flags", ["flag", "room", "surface", "rule id", "risk", "rule", "triggered by", "recommendation"], rows)

    rows = [[s["id"], s.get("room_id"), s.get("surface_id"), s.get("code"), s.get("description"),
             *_m(s.get("quantity")), s.get("unit"), s.get("source")] for s in res.get("scope", [])]
    _sheet(wb, "Scope", ["item", "room", "surface", "code", "description", "quantity", "quantity low",
                         "quantity high", "unit", "from"], rows)

    _sheet(wb, "Warnings", ["warning"], [[w] for w in res.get("warnings", [])])
    return wb


def write_sheet(res: dict, path: Path) -> Path:
    path = Path(path)
    build_workbook(res).save(path)
    return path


def main(argv: list[str] | None = None) -> None:
    import argparse
    ap = argparse.ArgumentParser(description="result.json -> result.xlsx")
    ap.add_argument("result", type=Path)
    ap.add_argument("--out", type=Path)
    a = ap.parse_args(argv)
    res = json.loads(a.result.read_text(encoding="utf-8"))
    print("wrote", write_sheet(res, a.out or a.result.with_suffix(".xlsx")))


if __name__ == "__main__":
    main()
